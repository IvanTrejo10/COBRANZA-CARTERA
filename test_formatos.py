# -*- coding: utf-8 -*-
"""Prueba funcional: datos sucios tipicos -> formato identico al BI."""
import io
import sys
import zipfile
from decimal import Decimal
import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, '/root/trabajo')
import datos_cobranza_latam as dl
import datos_cobranza_presico as dp
import procesos as pr

ok = True
def check(cond, msg):
    global ok
    print(('  OK   ' if cond else '  FALLA') + ' ' + msg)
    if not cond:
        ok = False

# ---------- convertir_numero_bi ----------
casos = [
    (Decimal('123.45'), 123.45), ('1,234.56', 1234.56), ('$ 2,500', 2500.0),
    ('(150.25)', -150.25), ('1.234.567', 1234567.0), ('12,5', 12.5),
    ('', None), ('NULL', None), (None, None), (b'99.9', 99.9),
    ('Q1,250.00', 1250.0), ('S/ 3,400.50', 3400.5), ('L 750', 750.0),
    ('C$ 1,000', 1000.0), (' 45 ', 45.0), ('-87.3', -87.3),
    (float('nan'), None), (0, 0.0), ('0.00', 0.0), ('1,234', 1234.0),
]
print('convertir_numero_bi:')
for entrada, esperado in casos:
    res = dl.convertir_numero_bi(entrada)
    check(res == esperado, f'{entrada!r} -> {res!r} (esperado {esperado!r})')

# ---------- limpiar_texto_bi / rutas ----------
print('normalizar rutas:')
check(dl.normalizar_ruta_latam('  ruta norte​ ') == 'RUTA NORTE', 'LATAM: espacios+minusculas+ZWSP -> RUTA NORTE')
check(dl.normalizar_ruta_latam(None) is None, 'LATAM: None se conserva')
check(dp.normalizar_ruta_presico_mx('MANZANILLO') == 'MANZANILLO-P-R2', 'PRESICO MX: MANZANILLO -> MANZANILLO-P-R2')
check(dp.normalizar_ruta_presico_mx(' MANZANILLO ') == 'MANZANILLO', 'PRESICO MX: reemplazo solo si valor exacto (igual que el BI)')
check(dp.normalizar_ruta_presico_mx('  COLIMA-R1 \x00') == 'COLIMA-R1', 'PRESICO MX: trim + clean')

# ---------- aplicar_formato_bi LATAM ----------
print('aplicar_formato_bi LATAM (datos crudos):')
df = pd.DataFrame({
    'Ruta': ['  ruta sur ', 'CENTRO​', None],
    'Zona': [b'ZONA A', 'ZONA B', 5],
    'coordinadora_id': ['1,234', 4567.0, None],
    'dbc_cliente_id': [Decimal('111'), None, '222'],
    'numero_clientes_dia': ['12', 15.0, None],
    'cobranza_del_dia': [Decimal('1234.50'), '2,345.75', 'NULL'],
    'saldo_devengado': ['(50.25)', 0, '$1,000'],
    '1_SEMANA_SALDO_DEVENGADO': ['10.5', Decimal('20.25'), ''],
    'CAMBIO_COORDINAODRA': [None, '2', 1.0],
    'TIPO_DESEMBOLSO': ['EFECTIVO', None, b'TRANSFERENCIA'],
})
r = dl.aplicar_formato_bi(df)
check(r['Ruta'][0] == 'RUTA SUR' and r['Ruta'][1] == 'CENTRO' and pd.isna(r['Ruta'][2]), 'Ruta normalizada como el BI')
check(r['coordinadora_id'].tolist() == [1234, 4567, pd.NA] or (r['coordinadora_id'][0] == 1234 and r['coordinadora_id'][1] == 4567 and pd.isna(r['coordinadora_id'][2])), 'coordinadora_id entera con nulos')
check(str(r['coordinadora_id'].dtype) == 'Int64', 'coordinadora_id dtype Int64 (sin .0)')
check(r['cobranza_del_dia'].tolist()[:2] == [1234.5, 2345.75] and pd.isna(r['cobranza_del_dia'][2]), 'cobranza_del_dia decimal desde Decimal/texto con comas')
check(r['saldo_devengado'].tolist()[:2] == [-50.25, 0.0] and r['saldo_devengado'][2] == 1000.0, 'saldo_devengado: parentesis negativo y $')
check(r['numero_clientes_dia'][0] == 12 and r['numero_clientes_dia'][1] == 15 and pd.isna(r['numero_clientes_dia'][2]), 'conteos enteros')
check(r['Zona'].tolist() == ['ZONA A', 'ZONA B', '5'], 'texto: bytes->str y numero->texto')
check(pd.isna(r['CAMBIO_COORDINAODRA'][0]) and r['CAMBIO_COORDINAODRA'][1] == 2, 'CAMBIO_COORDINAODRA entera nullable (como Int64 del BI)')

# ---------- etapa final LATAM (nombres renombrados) ----------
print('aplicar_formato_bi LATAM (nombres finales):')
df2 = pd.DataFrame({
    'ATRASO SEMANA 26-39': ['1,500.75', Decimal('10')],
    'PAGO COBRANZA EN ATRASO': ['200.10', '300'],
    'No. CLIENTES CON PAGO': ['8', 9.0],
    'dbc_cliente limpio': [123456.0, '789'],
    'TIPO CAMBIO.1': ['0.052', Decimal('0.13')],
})
r2 = dl.aplicar_formato_bi(df2)
check(r2['ATRASO SEMANA 26-39'].tolist() == [1500.75, 10.0], 'ATRASO SEMANA 26-39 numerica (corrige el tipo texto del BI)')
check(r2['No. CLIENTES CON PAGO'].tolist() == [8, 9], 'conteo final entero')
check(r2['dbc_cliente limpio'].tolist() == ['123456', '789'], 'dbc_cliente limpio texto sin .0 (como el BI LATAM)')
check(r2['TIPO CAMBIO.1'].tolist() == [0.052, 0.13], 'tipo de cambio decimal')

# ---------- aplicar_formato_bi PRESICO ----------
print('aplicar_formato_bi PRESICO:')
df3 = pd.DataFrame({
    'Ruta': ['MANZANILLO', ' tecoman-r2 '],
    'dbc_cliente limpio': ['1,001', 2002.0],
    'pago_cobranza_dia': [Decimal('500.25'), '1,250.50'],
    'numero_clientes_dia': ['20', None],
    'ENTREGADO_CLIENTE': ['$3,000', '(75.5)'],
    'Faltas': [3.0, '4'],
})
r3 = dp.aplicar_formato_bi(df3, servidor='PRESICO MX NV')
check(r3['Ruta'].tolist() == ['MANZANILLO-P-R2', 'tecoman-r2'], 'PRESICO MX NV: Ruta con reemplazo+trim (sin mayusculas, igual que el BI)')
check(r3['dbc_cliente limpio'].tolist() == [1001, 2002], 'dbc_cliente limpio ENTERO en Presico (como el BI)')
check(str(r3['dbc_cliente limpio'].dtype) == 'Int64', 'dtype Int64')
check(r3['pago_cobranza_dia'].tolist() == [500.25, 1250.5], 'dinero decimal')
check(r3['ENTREGADO_CLIENTE'].tolist() == [3000.0, -75.5], 'ENTREGADO_CLIENTE con $ y negativo')
check(r3['Faltas'].tolist() == [3, 4], 'Faltas entera')
r3b = dp.aplicar_formato_bi(df3, servidor='LA CASITA')
check(r3b['Ruta'].tolist() == ['MANZANILLO', ' tecoman-r2 '], 'otros servidores: Ruta intacta (igual que el BI)')

# ---------- redondeo opcional ----------
r4 = dp.aplicar_formato_bi(pd.DataFrame({'pago_servicio': [10.005, '33.333']}), redondear=2)
check(r4['pago_servicio'].tolist() == [10.0, 33.33] or r4['pago_servicio'].tolist() == [10.01, 33.33], 'redondear=2 opcional')

# ---------- exportacion Excel: Grupo siempre como texto ----------
print('excel_bytes Grupo como texto:')
df_grupos = pd.DataFrame({
    'Grupo': [119, 120.0, '00121', None],
    'Monto': [10.0, 20.0, 30.0, 40.0],
})
excel_grupos = pr.excel_bytes(df_grupos)
wb = load_workbook(io.BytesIO(excel_grupos), data_only=True)
ws = wb['Hoja1']
celdas_grupo = [ws.cell(row=fila, column=1) for fila in range(2, 6)]
check([c.value for c in celdas_grupo] == ['119', '120', '00121', None],
      'valores de Grupo conservados como texto, incluidos ceros a la izquierda')
check(all(c.data_type == 's' and c.number_format == '@' for c in celdas_grupo[:3]),
      'celdas no vacias de Grupo tienen tipo string y formato Texto (@)')
check(celdas_grupo[3].number_format == '@', 'celdas vacias de Grupo conservan formato Texto (@)')
with zipfile.ZipFile(io.BytesIO(excel_grupos)) as archivo_xlsx:
    xml_hoja = archivo_xlsx.read('xl/worksheets/sheet1.xml').decode('utf-8')
check('sqref="A2:A5"' in xml_hoja and 'numberStoredAsText="1"' in xml_hoja,
      'Excel no muestra la advertencia de numero almacenado como texto en Grupo')

# ---------- las variables originales siguen intactas ----------
print('integridad del modulo:')
check(len(dl.RAW_QUERIES) == 10 and len(dp.RAW_QUERIES_PRESICO) == 7, 'RAW_QUERIES completos')
check(dl.SERVIDORES_COBRANZA[0]['db'] == 'peru_presico' and dp.SERVIDORES_PRESICO[0]['db'] == 'ciudadespresico_newSystem_mx', 'servidores intactos')
check('#(lf)' in dl.RAW_QUERIES['Cobranza Colombia AWS'], 'queries tal cual (escapes de Power Query conservados)')

print()
print('RESULTADO FINAL:', 'TODO OK' if ok else 'HAY FALLAS')
sys.exit(0 if ok else 1)
